from openpyxl.workbook import  Workbook
from openpyxl.chart import Series,LineChart,Reference

wb = Workbook()
ws = wb.create_sheet("{0}蒸汽压力记录表".format("样品名称"),0)

chart = LineChart()
data = Reference(ws,min_col=5,min_row=4,max_col=10,max_row=4)
seriesObj = Series(data,title='压力')
chart.append(seriesObj)
ws.add_chart(chart,"A6")
wb.save("D://bar2.xlsx")